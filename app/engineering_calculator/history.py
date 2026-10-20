import csv
import io
import json
import sqlite3
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DATABASE_PATH = Path("data/calculation_history.db")


def initialize_database():
    DATABASE_PATH.parent.mkdir(parents=True, exist_ok=True)

    with sqlite3.connect(DATABASE_PATH) as connection:
        connection.execute("""
            CREATE TABLE IF NOT EXISTS calculations (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT NOT NULL,
                formula TEXT NOT NULL,
                inputs TEXT NOT NULL,
                result REAL NOT NULL,
                unit TEXT NOT NULL
            )
        """)


def save_calculation(formula, inputs, result, unit):
    initialize_database()

    with sqlite3.connect(DATABASE_PATH) as connection:
        connection.execute(
            """
            INSERT INTO calculations
                (created_at, formula, inputs, result, unit)
            VALUES (?, ?, ?, ?, ?)
        """,
            (
                datetime.now(UTC).isoformat(timespec="seconds"),
                formula,
                json.dumps(inputs),
                float(result),
                unit,
            ),
        )


def get_history(limit=100):
    initialize_database()

    with sqlite3.connect(DATABASE_PATH) as connection:
        connection.row_factory = sqlite3.Row

        rows = connection.execute(
            """
            SELECT id, created_at, formula, inputs, result, unit
            FROM calculations
            ORDER BY id DESC
            LIMIT ?
        """,
            (max(1, min(int(limit), 1000)),),
        ).fetchall()

    history = []

    for row in rows:
        item = dict(row)
        item["inputs"] = json.loads(item["inputs"])
        history.append(item)

    return history


def export_history_csv():
    records = get_history(limit=1000)
    output = io.StringIO()

    writer = csv.writer(output)
    writer.writerow(
        [
            "ID",
            "Timestamp (UTC)",
            "Formula",
            "Inputs",
            "Result",
            "Unit",
        ]
    )

    for record in reversed(records):
        writer.writerow(
            [
                record["id"],
                record["created_at"],
                record["formula"],
                json.dumps(record["inputs"]),
                record["result"],
                record["unit"],
            ]
        )

    return output.getvalue()

def export_history_excel():
    records = get_history(limit=1000)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Calculation History"

    headers = [
        "ID",
        "Timestamp (UTC)",
        "Formula",
        "Inputs",
        "Result",
        "Unit",
    ]

    worksheet.append(headers)

    for record in reversed(records):
        worksheet.append(
            [
                record["id"],
                record["created_at"],
                record["formula"],
                json.dumps(
                    record["inputs"],
                    ensure_ascii=False,
                ),
                record["result"],
                record["unit"],
            ]
        )

    # Professional header formatting
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="123047",
    )
    header_font = Font(
        bold=True,
        color="FFFFFF",
        size=11,
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # Automatically adjust column widths
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(
            column_cells[0].column
        )

        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )

        worksheet.column_dimensions[
            column_letter
        ].width = min(max(max_length + 2, 12), 45)

    # Make long input data readable
    for row in worksheet.iter_rows(min_row=2):
        row[3].alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 25

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output.getvalue()
